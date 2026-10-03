"""Ноутбук должен быть синтаксически исполним.

Ловит регрессию, из-за которой ячейка 10 падала прямо в Colab с
`SyntaxError: unterminated f-string literal` (генератор подставил реальные переводы
строки внутрь f-строки).
"""

from __future__ import annotations

import ast
import json
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
NOTEBOOK = ROOT / "BTZSC_Colab.ipynb"

# В отдельном bundle ноутбука нет, поэтому внутри Colab эти проверки пропускаются.
pytestmark = pytest.mark.skipif(not NOTEBOOK.exists(), reason="BTZSC_Colab.ipynb недоступен (запуск из bundle)")


def load_cells() -> list[dict]:
    if not NOTEBOOK.exists():
        return []
    return json.loads(NOTEBOOK.read_text(encoding="utf-8"))["cells"]


def as_python(source: str) -> str:
    """IPython-магии (%cd) и shell-команды (!pip) в обычный Python не парсятся."""
    lines = []
    for line in source.split("\n"):
        stripped = line.strip()
        lines.append("pass  # " + line if stripped.startswith(("!", "%")) else line)
    return "\n".join(lines)


@pytest.mark.parametrize("index", range(max(1, len(load_cells()))))
def test_code_cell_compiles(index: int):
    cells = load_cells()
    if index >= len(cells):
        return
    cell = cells[index]
    if cell["cell_type"] != "code":
        return
    source = "".join(cell["source"])
    try:
        ast.parse(as_python(source))
    except SyntaxError as err:  # pragma: no cover - сообщение важнее покрытия
        pytest.fail(f"ячейка {index}: {err.msg} (строка {err.lineno})\n{source[:400]}")


def test_notebook_calls_experiment_module():
    text = "".join("".join(c["source"]) for c in load_cells())
    assert "import experiment" in text, "ноутбук должен вызывать experiment.py, а не дублировать логику"


def test_notebook_has_no_embedded_archive():
    text = "".join("".join(c["source"]) for c in load_cells())
    assert "BUNDLE_B64" not in text
    assert "BUNDLE_PAYLOAD" not in text


def test_no_secrets_hardcoded():
    text = "".join("".join(c["source"]) for c in load_cells())
    for marker in ("hf_", "BEGIN OPENSSH PRIVATE KEY", "BEGIN RSA PRIVATE KEY", "ghp_"):
        assert marker not in text.replace("hf_cache", "").replace("HF_TOKEN", ""), f"в ноутбуке найдено: {marker}"


def test_code_comes_from_public_repo():
    text = "".join("".join(c["source"]) for c in load_cells())
    assert "https://github.com/clinch0/ML2" in text
    assert '"clone"' in text and '"--depth"' in text
    assert "BUNDLE_B64" not in text


def test_presentation_builds_with_speech_and_figures():
    """Колода: 12 слайдов, у каждого речь, фигуры-объяснялки вставлены."""
    import subprocess
    import sys
    from pathlib import Path

    root = Path(__file__).resolve().parent.parent
    out = root / "Задание2_BTZSC.pptx"
    proc = subprocess.run([sys.executable, "scripts/build_btzsc_pptx.py", "--out", str(out)],
                          cwd=root, capture_output=True, text=True, timeout=300)
    assert proc.returncode == 0, proc.stdout + proc.stderr

    from pptx import Presentation

    prs = Presentation(str(out))
    assert len(prs.slides) == 12

    words = 0
    for i, slide in enumerate(prs.slides, 1):
        assert slide.has_notes_slide, f"слайд {i} без речи докладчика"
        text = slide.notes_slide.notes_text_frame.text.strip()
        assert len(text.split()) >= 25, f"слайд {i}: речь слишком короткая"
        words += len(text.split())
    assert 550 <= words <= 650, f"речь должна занимать ~4 минуты, сейчас {words} слов"

    pictures = sum(1 for slide in prs.slides for sh in slide.shapes if sh.shape_type == 13)
    assert pictures >= 5, f"фигур на слайдах мало: {pictures}"


def test_slide_figures_are_opaque():
    """Фигуры сохраняются на светлом непрозрачном фоне: в тёмной теме иначе пропадает текст."""
    from pathlib import Path

    src = (Path(__file__).resolve().parent.parent / "scripts" / "slide_figures.py").read_text(encoding="utf-8")
    assert 'facecolor="white"' in src and "transparent=True" not in src

    main_src = (Path(__file__).resolve().parent.parent / "scripts" / "make_figures.py").read_text(encoding="utf-8")
    assert "transparent=True" not in main_src, "прозрачный фон убран и в основном скрипте"
